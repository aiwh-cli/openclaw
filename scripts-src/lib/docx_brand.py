"""
AIWH Branded DOCX Builder — Light Theme
Generates styled .docx files matching aiwh-website light theme.
Uses Arial as Google Docs-compatible fallback for Inter.

Usage:
    from docx_brand import BrandDoc
    doc = BrandDoc("Document Title")
    doc.h1("Section Heading")
    doc.body("Paragraph text here.")
    doc.bullet("First bullet point")
    doc.callout("Important note in a callout box")
    doc.table([["Col A","Col B"],["val1","val2"]])
    doc.save("/path/to/output.docx")
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import re

# ── Brand Colors (Light Theme) ──────────────────────────────────
ACCENT        = RGBColor(0x9B, 0x7D, 0x2E)  # #9B7D2E
ACCENT_HOT    = RGBColor(0xC9, 0xA8, 0x4C)  # #C9A84C
TEXT_PRIMARY   = RGBColor(0x1A, 0x18, 0x15)  # #1A1815
TEXT_MUTED     = RGBColor(0x6B, 0x65, 0x60)  # #6B6560
TEXT_DIM       = RGBColor(0xA0, 0x9A, 0x90)  # #A09A90
BG_WARM        = RGBColor(0xFA, 0xF8, 0xF3)  # #FAF8F3
WHITE          = RGBColor(0xFF, 0xFF, 0xFF)

FONT_BODY     = "Arial"       # Google Docs compatible fallback for Inter
FONT_MONO     = "Courier New" # Google Docs compatible fallback for JetBrains Mono


class BrandDoc:
    def __init__(self, title, subtitle=None):
        self.doc = Document()
        self._setup_styles()
        self._setup_margins()
        self._add_title_block(title, subtitle)

    def _setup_margins(self):
        for section in self.doc.sections:
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(2.54)
            section.right_margin = Cm(2.54)

    def _setup_styles(self):
        styles = self.doc.styles

        # Normal body
        normal = styles['Normal']
        normal.font.name = FONT_BODY
        normal.font.size = Pt(11)
        normal.font.color.rgb = TEXT_PRIMARY
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15

        # Title
        title_style = styles['Title']
        title_style.font.name = FONT_BODY
        title_style.font.size = Pt(28)
        title_style.font.bold = True
        title_style.font.color.rgb = TEXT_PRIMARY
        title_style.paragraph_format.space_after = Pt(4)

        # Subtitle
        subtitle_style = styles['Subtitle']
        subtitle_style.font.name = FONT_BODY
        subtitle_style.font.size = Pt(13)
        subtitle_style.font.bold = False
        subtitle_style.font.color.rgb = TEXT_MUTED
        subtitle_style.paragraph_format.space_after = Pt(18)

        # Heading 1
        h1 = styles['Heading 1']
        h1.font.name = FONT_BODY
        h1.font.size = Pt(18)
        h1.font.bold = True
        h1.font.color.rgb = TEXT_PRIMARY
        h1.paragraph_format.space_before = Pt(24)
        h1.paragraph_format.space_after = Pt(8)

        # Heading 2 — accent color
        h2 = styles['Heading 2']
        h2.font.name = FONT_BODY
        h2.font.size = Pt(14)
        h2.font.bold = True
        h2.font.color.rgb = ACCENT
        h2.paragraph_format.space_before = Pt(18)
        h2.paragraph_format.space_after = Pt(6)

        # Heading 3 — muted
        h3 = styles['Heading 3']
        h3.font.name = FONT_BODY
        h3.font.size = Pt(12)
        h3.font.bold = True
        h3.font.color.rgb = TEXT_MUTED
        h3.paragraph_format.space_before = Pt(14)
        h3.paragraph_format.space_after = Pt(4)

        # List Bullet
        lb = styles['List Bullet']
        lb.font.name = FONT_BODY
        lb.font.size = Pt(11)
        lb.font.color.rgb = TEXT_PRIMARY
        lb.paragraph_format.space_after = Pt(3)

    def _add_title_block(self, title, subtitle):
        # Gold accent line before title
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run("━" * 40)
        run.font.color.rgb = ACCENT
        run.font.size = Pt(8)

        self.doc.add_paragraph(title, style='Title')
        if subtitle:
            self.doc.add_paragraph(subtitle, style='Subtitle')

        # Divider after title
        self._add_divider()

    def _add_divider(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run("━" * 60)
        run.font.color.rgb = ACCENT_HOT
        run.font.size = Pt(6)

    def h1(self, text):
        self.doc.add_paragraph(text, style='Heading 1')

    def h2(self, text):
        self.doc.add_paragraph(text, style='Heading 2')

    def h3(self, text):
        self.doc.add_paragraph(text, style='Heading 3')

    def body(self, text):
        """Add a body paragraph. Supports **bold** and *italic* inline."""
        p = self.doc.add_paragraph(style='Normal')
        self._add_formatted_runs(p, text)

    def bold_body(self, text):
        p = self.doc.add_paragraph(style='Normal')
        run = p.add_run(text)
        run.bold = True

    def bullet(self, text):
        p = self.doc.add_paragraph(style='List Bullet')
        self._add_formatted_runs(p, text)

    def numbered(self, text):
        p = self.doc.add_paragraph(style='List Number')
        self._add_formatted_runs(p, text)

    def callout(self, text, label="NOTE"):
        """Callout box with warm background and accent left border."""
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6)

        # Background color
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAF8F3" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shading)

        # Left border accent
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="9B7D2E"/>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E8E4DC"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E8E4DC"/>'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="E8E4DC"/>'
            f'</w:tcBorders>'
        )
        tc_pr.append(borders)

        p = cell.paragraphs[0]
        label_run = p.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.color.rgb = ACCENT
        label_run.font.size = Pt(10)
        label_run.font.name = FONT_BODY

        text_run = p.add_run(text)
        text_run.font.color.rgb = TEXT_PRIMARY
        text_run.font.size = Pt(10)
        text_run.font.name = FONT_BODY

        # Spacing after callout
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def code(self, text):
        """Code block with dark background."""
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)

        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1A1815" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shading)

        for i, line in enumerate(text.strip().split('\n')):
            if i == 0:
                p = cell.paragraphs[0]
            else:
                p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            run = p.add_run(line)
            run.font.name = FONT_MONO
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xF5, 0xED, 0xD6)  # #F5EDD6

        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, rows, col_widths=None):
        """Branded table. First row = header with accent background."""
        if not rows or len(rows) < 1:
            return

        num_cols = len(rows[0])
        tbl = self.doc.add_table(rows=len(rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.style = 'Table Grid'

        for i, row_data in enumerate(rows):
            for j, cell_text in enumerate(row_data):
                cell = tbl.cell(i, j)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                if i == 0:
                    # Header row — accent background, white text
                    shading = parse_xml(
                        f'<w:shd {nsdecls("w")} w:fill="9B7D2E" w:val="clear"/>'
                    )
                    cell._tc.get_or_add_tcPr().append(shading)
                    run = p.add_run(str(cell_text))
                    run.bold = True
                    run.font.color.rgb = WHITE
                    run.font.size = Pt(10)
                    run.font.name = FONT_BODY
                else:
                    # Alt row shading
                    if i % 2 == 0:
                        shading = parse_xml(
                            f'<w:shd {nsdecls("w")} w:fill="FAF8F3" w:val="clear"/>'
                        )
                        cell._tc.get_or_add_tcPr().append(shading)
                    run = p.add_run(str(cell_text))
                    run.font.size = Pt(10)
                    run.font.name = FONT_BODY
                    run.font.color.rgb = TEXT_PRIMARY

        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def spacer(self, pts=12):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(pts)

    def divider(self):
        self._add_divider()

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)
        return path

    def _add_formatted_runs(self, paragraph, text):
        """Parse **bold** and *italic* markers in text."""
        parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                paragraph.add_run(part)
